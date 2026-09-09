import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox

from openpyxl import load_workbook
import win32com.client as win32


VERSION = "v4"


def choose_excel():
    path = filedialog.askopenfilename(
        title="Επιλογή αρχείου Excel",
        filetypes=[("Excel files", "*.xlsx")]
    )
    if path:
        excel_var.set(path)


def choose_template():
    path = filedialog.askopenfilename(
        title="Επιλογή προτύπου Word",
        filetypes=[("Word documents", "*.docx")]
    )
    if path:
        template_var.set(path)


def choose_output():
    path = filedialog.askdirectory(
        title="Επιλογή φακέλου αποθήκευσης"
    )
    if path:
        output_var.set(path)


def clean_value(value):
    if value is None:
        return ""

    if hasattr(value, "strftime"):
        return value.strftime("%d/%m/%Y")

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value)


def safe_filename(text):
    return re.sub(r'[<>:"/\\|?*]', "_", text)


def replace_exact_in_range(word_range, placeholder, value):
    """
    Βρίσκει το placeholder και αντικαθιστά ΜΟΝΟ
    το συγκεκριμένο εύρος χαρακτήρων.
    """

    search_range = word_range.Duplicate

    while True:
        find = search_range.Find

        find.ClearFormatting()
        find.Text = placeholder
        find.Forward = True
        find.Wrap = 0
        find.Format = False
        find.MatchCase = True
        find.MatchWholeWord = False
        find.MatchWildcards = False

        found = find.Execute()

        if not found:
            break

        # Μετά το Find, το search_range είναι ακριβώς
        # πάνω στο κείμενο που βρέθηκε.
        search_range.Text = str(value)

        # Συνεχίζουμε την αναζήτηση μετά την αντικατάσταση
        new_start = search_range.End

        search_range.SetRange(
            Start=new_start,
            End=word_range.End
        )


def replace_word_text(doc, placeholder, value):

    # Κυρίως σώμα
    replace_exact_in_range(
        doc.Content,
        placeholder,
        value
    )

    # Headers / Footers
    for section in doc.Sections:

        for header in section.Headers:
            if header.Exists:
                replace_exact_in_range(
                    header.Range,
                    placeholder,
                    value
                )

        for footer in section.Footers:
            if footer.Exists:
                replace_exact_in_range(
                    footer.Range,
                    placeholder,
                    value
                )


def create_documents():

    excel_path = excel_var.get()
    template_path = template_var.get()
    output_folder = output_var.get()

    if not excel_path or not template_path or not output_folder:
        messagebox.showwarning(
            "Λείπουν στοιχεία",
            "Επίλεξε Excel, Word Template και φάκελο αποθήκευσης."
        )
        return

    word = None

    try:

        wb = load_workbook(
            excel_path,
            data_only=True
        )

        ws = wb.active

        headers = [
            clean_value(cell.value).strip()
            for cell in ws[1]
        ]

        if not any(headers):
            raise Exception(
                "Δεν βρέθηκαν επικεφαλίδες "
                "στην πρώτη γραμμή του Excel."
            )

        word = win32.DispatchEx(
            "Word.Application"
        )

        word.Visible = False
        word.DisplayAlerts = False

        created = 0

        for row_number, row in enumerate(
            ws.iter_rows(
                min_row=2,
                values_only=True
            ),
            start=2
        ):

            if all(value is None for value in row):
                continue

            data = {}

            for i in range(
                min(len(headers), len(row))
            ):

                header = headers[i]

                if not header:
                    continue

                data[header] = clean_value(
                    row[i]
                ).strip()

            doc = word.Documents.Open(
                os.path.abspath(template_path),
                ReadOnly=False
            )

            # ---------------------------------------
            # ΑΝΤΙΚΑΤΑΣΤΑΣΗ PLACEHOLDERS
            # ---------------------------------------

            for header, value in data.items():

                placeholder = f"[{header}]"

                replace_word_text(
                    doc,
                    placeholder,
                    value
                )

            # ---------------------------------------
            # ΕΛΕΓΧΟΣ
            # ---------------------------------------

            remaining = []

            for header in data:

                placeholder = f"[{header}]"

                if placeholder in doc.Content.Text:
                    remaining.append(placeholder)

            # ---------------------------------------
            # ΟΝΟΜΑ ΑΡΧΕΙΟΥ
            # ---------------------------------------

            surname = data.get(
                "ΕΠΩΝΥΜΟ",
                ""
            )

            name = data.get(
                "ΟΝΟΜΑ",
                ""
            )

            am = data.get(
                "ΑΜ",
                ""
            )

            if surname or name:

                filename = (
                    f"{VERSION}_"
                    f"{surname}_{name}"
                )

                if am:
                    filename += f"_{am}"

            else:

                filename = (
                    f"{VERSION}_"
                    f"ΕΓΓΡΑΦΟ_{row_number}"
                )

            filename = safe_filename(
                filename
            ) + ".docx"

            output_path = os.path.abspath(
                os.path.join(
                    output_folder,
                    filename
                )
            )

            if os.path.exists(output_path):
                os.remove(output_path)

            doc.SaveAs2(
                output_path,
                FileFormat=16
            )

            doc.Close(
                SaveChanges=False
            )

            created += 1

        word.Quit()
        word = None

        messagebox.showinfo(
            "Ολοκληρώθηκε",
            f"Έκδοση {VERSION}\n\n"
            f"Δημιουργήθηκαν {created} έγγραφα."
        )

    except Exception as e:

        if word is not None:
            try:
                word.Quit()
            except:
                pass

        messagebox.showerror(
            "Σφάλμα",
            f"Παρουσιάστηκε σφάλμα:\n\n{e}"
        )


# ============================================================
# GUI
# ============================================================

root = tk.Tk()

root.title(
    f"Excel → Word Merge {VERSION}"
)

root.geometry(
    "700x330"
)

root.resizable(
    False,
    False
)

excel_var = tk.StringVar()
template_var = tk.StringVar()
output_var = tk.StringVar()


tk.Label(
    root,
    text=f"Δημιουργία ατομικών εγγράφων — {VERSION}",
    font=("Arial", 16, "bold")
).pack(
    pady=18
)


frame = tk.Frame(root)

frame.pack(
    fill="x",
    padx=20
)


tk.Button(
    frame,
    text="Επιλογή Excel",
    width=22,
    command=choose_excel
).grid(
    row=0,
    column=0,
    pady=8
)

tk.Entry(
    frame,
    textvariable=excel_var,
    width=60
).grid(
    row=0,
    column=1,
    padx=10
)


tk.Button(
    frame,
    text="Επιλογή Word Template",
    width=22,
    command=choose_template
).grid(
    row=1,
    column=0,
    pady=8
)

tk.Entry(
    frame,
    textvariable=template_var,
    width=60
).grid(
    row=1,
    column=1,
    padx=10
)


tk.Button(
    frame,
    text="Φάκελος αποθήκευσης",
    width=22,
    command=choose_output
).grid(
    row=2,
    column=0,
    pady=8
)

tk.Entry(
    frame,
    textvariable=output_var,
    width=60
).grid(
    row=2,
    column=1,
    padx=10
)


tk.Button(
    root,
    text="ΔΗΜΙΟΥΡΓΙΑ ΕΓΓΡΑΦΩΝ",
    font=("Arial", 12, "bold"),
    width=30,
    height=2,
    command=create_documents
).pack(
    pady=25
)


root.mainloop()
